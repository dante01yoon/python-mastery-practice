from openpyxl import Workbook, load_workbook

wb = Workbook()
ws = wb.active
ws.title = "매출"
ws.append(["날짜", "상품", "금액"])
ws.append(["10-27", "커피", 4500])
ws.append(["10-28", "케이크", 6800])
ws.append(["10-29", "커피", 4500])
ws["E1"] = "합계"
ws["F1"] = "=SUM(C2:C4)"
wb.save("sales.xlsx")

wb2 = load_workbook("sales.xlsx")
total = 0
for row in wb2["매출"].iter_rows(min_row=2, values_only=True):
    print(row[:3])
    total += row[2]
print("합계:", total)
